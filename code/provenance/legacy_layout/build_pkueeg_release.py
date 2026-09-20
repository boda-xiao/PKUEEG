#!/usr/bin/env python3
"""Build a non-destructive PKUEEG release candidate from existing server data.

The script never edits the source tree. Large unchanged files are hard-linked when
source and output are on the same filesystem. Files that require correction are
written to temporary files and atomically moved into the release tree so source
hard links are not modified.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import shutil
import subprocess
from concurrent.futures import ProcessPoolExecutor, as_completed
from pathlib import Path

import numpy as np
from scipy.signal import butter, resample_poly, sosfiltfilt


SUBJECTS = [f"sub-{index:02d}" for index in range(1, 26)]
SESSION_STORIES = {
    "day1": list(range(1, 18)),
    "day2": list(range(18, 34)),
    "day3": list(range(34, 51)),
}
AUTHORS = [
    "Boda Xiao",
    "X. Xu",
    "Y. Yan",
    "L. Zheng",
    "S. Li",
    "Z. Zhang",
    "J. Liang",
    "A. Zhang",
    "Xihong Wu",
    "Heping Cheng",
    "Jing Chen",
]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-root", type=Path, required=True)
    parser.add_argument("--output-root", type=Path, required=True)
    parser.add_argument("--behavior-xlsx", type=Path, required=True)
    parser.add_argument("--stimulus-manifest", type=Path)
    parser.add_argument("--workers", type=int, default=4)
    parser.add_argument(
        "--resume",
        action="store_true",
        help="Resume an existing candidate after verifying that Day-3 raw scaling is already complete.",
    )
    return parser.parse_args()


def atomic_write_text(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(f".{path.name}.tmp-{os.getpid()}")
    tmp.write_text(text, encoding="utf-8")
    os.replace(tmp, path)


def atomic_write_json(path: Path, value: object) -> None:
    atomic_write_text(path, json.dumps(value, ensure_ascii=False, indent=2) + "\n")


def atomic_write_tsv(path: Path, fieldnames: list[str], rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(f".{path.name}.tmp-{os.getpid()}")
    with tmp.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, delimiter="\t", lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)
    os.replace(tmp, path)


def hardlink_or_copy(source: Path, target: Path) -> None:
    target.parent.mkdir(parents=True, exist_ok=True)
    try:
        os.link(source, target)
    except OSError:
        shutil.copy2(source, target)


def clone_raw_bids(source: Path, target: Path) -> None:
    target.mkdir(parents=True)
    subprocess.run(["cp", "-al", f"{source}/.", str(target)], check=True)


def scale_float32_file_atomic(path: Path, factor: float, block_values: int = 16_777_216) -> None:
    size = path.stat().st_size
    if size % np.dtype("<f4").itemsize:
        raise ValueError(f"Float32 file has invalid size: {path}")
    source = np.memmap(path, dtype="<f4", mode="r")
    tmp = path.with_name(f".{path.name}.scaled-{os.getpid()}")
    target = np.memmap(tmp, dtype="<f4", mode="w+", shape=source.shape)
    for start in range(0, source.size, block_values):
        stop = min(start + block_values, source.size)
        target[start:stop] = source[start:stop] * factor
    target.flush()
    del target
    del source
    shutil.copystat(path, tmp)
    os.replace(tmp, path)


def remove_unmeasured_coordinates(root: Path) -> int:
    files = list(root.glob("sub-*/ses-*/eeg/*_electrodes.tsv"))
    files += list(root.glob("sub-*/ses-*/eeg/*_coordsystem.json"))
    for path in files:
        path.unlink()
    return len(files)


def update_raw_metadata(root: Path) -> None:
    for subject in SUBJECTS:
        for day in (1, 2, 3):
            session = f"day{day}"
            eeg_dir = root / subject / f"ses-{session}" / "eeg"
            stem = f"{subject}_ses-{session}_task-audio"
            sidecar_path = eeg_dir / f"{stem}_eeg.json"
            sidecar = json.loads(sidecar_path.read_text(encoding="utf-8-sig"))
            sidecar.update(
                {
                    "TaskName": "audio",
                    "TaskDescription": "Eyes-open rest followed by continuous Mandarin speech perception.",
                    "Instructions": "Fixate the central cross, listen attentively, and answer comprehension questions after each story.",
                    "InstitutionName": "Peking University",
                    "PowerLineFrequency": 50,
                    "EEGPlacementScheme": "International 10-20 system",
                }
            )
            if day in (1, 2):
                sidecar.update(
                    {
                        "Manufacturer": "Compumedics Neuroscan",
                        "ManufacturersModelName": "SynAmps2",
                        "SoftwareVersions": "Scan; version not recorded",
                    }
                )
            else:
                sidecar.update(
                    {
                        "Manufacturer": "Neuracle Medical Technology Co., Ltd.",
                        "ManufacturersModelName": "NeuSen Wireless EEG/ERP",
                        "SoftwareVersions": "NeuSen Recorder; version not recorded",
                    }
                )
            atomic_write_json(sidecar_path, sidecar)

            channels_path = eeg_dir / f"{stem}_channels.tsv"
            with channels_path.open("r", encoding="utf-8-sig", newline="") as handle:
                reader = csv.DictReader(handle, delimiter="\t")
                rows = list(reader)
                fields = list(reader.fieldnames or [])
            for row in rows:
                row["units"] = "V"
                row["low_cutoff"] = "n/a"
                row["high_cutoff"] = "n/a"
            atomic_write_tsv(channels_path, fields, rows)

            scans_path = root / subject / f"ses-{session}" / f"{subject}_ses-{session}_scans.tsv"
            with scans_path.open("r", encoding="utf-8-sig", newline="") as handle:
                reader = csv.DictReader(handle, delimiter="\t")
                rows = list(reader)
                fields = list(reader.fieldnames or [])
            for row in rows:
                if "acq_time" in row:
                    row["acq_time"] = "n/a"
            atomic_write_tsv(scans_path, fields, rows)


def write_participants(root: Path) -> None:
    rows = [
        {"participant_id": subject, "age": "n/a", "sex": "n/a", "hand": "R"}
        for subject in SUBJECTS
    ]
    atomic_write_tsv(root / "participants.tsv", ["participant_id", "age", "sex", "hand"], rows)
    atomic_write_json(
        root / "participants.json",
        {
            "participant_id": {"Description": "Anonymous participant identifier."},
            "age": {
                "Description": "Age is withheld at participant level; the cohort range was 19-29 years.",
                "Units": "years",
            },
            "sex": {
                "Description": "Sex is withheld at participant level; the cohort comprised 14 men and 11 women.",
                "Levels": {"F": "female", "M": "male"},
            },
            "hand": {"Description": "Self-reported handedness.", "Levels": {"R": "right-handed"}},
        },
    )


def write_sessions(root: Path) -> None:
    fields = [
        "session_id",
        "acquisition_day",
        "story_first",
        "story_last",
        "n_story_runs",
        "rest_duration_sec",
        "sampling_frequency",
        "acquisition_system",
        "vendor_source_format",
        "recorded_channels",
    ]
    rows = [
        {
            "session_id": "ses-day1",
            "acquisition_day": 1,
            "story_first": 1,
            "story_last": 17,
            "n_story_runs": 17,
            "rest_duration_sec": 300,
            "sampling_frequency": 1000,
            "acquisition_system": "NeuroScan SynAmps2",
            "vendor_source_format": "CNT",
            "recorded_channels": "62 EEG + 2 EOG",
        },
        {
            "session_id": "ses-day2",
            "acquisition_day": 2,
            "story_first": 18,
            "story_last": 33,
            "n_story_runs": 16,
            "rest_duration_sec": 300,
            "sampling_frequency": 1000,
            "acquisition_system": "NeuroScan SynAmps2",
            "vendor_source_format": "CNT",
            "recorded_channels": "62 EEG + 2 EOG",
        },
        {
            "session_id": "ses-day3",
            "acquisition_day": 3,
            "story_first": 34,
            "story_last": 50,
            "n_story_runs": 17,
            "rest_duration_sec": 300,
            "sampling_frequency": 1000,
            "acquisition_system": "NeuSen Wireless EEG/ERP",
            "vendor_source_format": "BDF",
            "recorded_channels": "59 EEG + 5 auxiliary",
        },
    ]
    schema = {
        "session_id": {"Description": "BIDS session identifier."},
        "acquisition_day": {"Description": "Ordinal recording day."},
        "story_first": {"Description": "First story identifier scheduled in the session."},
        "story_last": {"Description": "Last story identifier scheduled in the session."},
        "n_story_runs": {"Description": "Number of story recordings in the session."},
        "rest_duration_sec": {"Description": "Planned eyes-open resting-state duration.", "Units": "s"},
        "sampling_frequency": {"Description": "Acquisition sampling frequency.", "Units": "Hz"},
        "acquisition_system": {"Description": "EEG acquisition system."},
        "vendor_source_format": {"Description": "Original vendor file format."},
        "recorded_channels": {"Description": "Recorded scalp and auxiliary channel counts."},
    }
    obsolete_root_sidecar = root / "sessions.json"
    if obsolete_root_sidecar.exists():
        obsolete_root_sidecar.unlink()
    for subject in SUBJECTS:
        atomic_write_tsv(root / subject / f"{subject}_sessions.tsv", fields, rows)
        atomic_write_json(root / subject / f"{subject}_sessions.json", schema)


def replace_events_with_intervals(source_root: Path, release_root: Path) -> None:
    day12_trigger_root = source_root / "derivatives" / "formal_filer_40hz"
    day3_trigger_root = (
        source_root
        / "bids"
        / "derivatives"
        / "derivatives_npz"
        / "formal_filter_40hz"
    )
    fields = [
        "onset",
        "duration",
        "sample",
        "trial_type",
        "story_id",
        "stim_file",
        "candidate_pair_count",
        "selection_rule",
    ]
    for subject in SUBJECTS:
        for day, expected_stories in SESSION_STORIES.items():
            trigger_root = day3_trigger_root if day == "day3" else day12_trigger_root
            trigger_path = trigger_root / subject / f"trigger_info_{day}.csv"
            with trigger_path.open("r", encoding="utf-8-sig", newline="") as handle:
                source_rows = list(csv.DictReader(handle))
            by_id = {int(float(row["event_id"])): row for row in source_rows}
            expected_ids = set(expected_stories) | {255}
            if set(by_id) != expected_ids:
                raise ValueError(
                    f"Unexpected selected event IDs in {trigger_path}: "
                    f"missing={sorted(expected_ids - set(by_id))}, extra={sorted(set(by_id) - expected_ids)}"
                )
            fallback_pairs: dict[int, list[tuple[float, float]]] = {}
            if any(not row.get("all_pairs_count") for row in source_rows):
                session = f"ses-{day}"
                original_events = (
                    source_root
                    / "bids"
                    / "rawdata_bids"
                    / subject
                    / session
                    / "eeg"
                    / f"{subject}_{session}_task-audio_events.tsv"
                )
                with original_events.open("r", encoding="utf-8-sig", newline="") as handle:
                    marker_rows = list(csv.DictReader(handle, delimiter="\t"))
                for first, second in zip(marker_rows, marker_rows[1:]):
                    first_id = int(float(first["trial_type"]))
                    second_id = int(float(second["trial_type"]))
                    gap = float(second["onset"]) - float(first["onset"])
                    if first_id == second_id and gap >= 150:
                        fallback_pairs.setdefault(first_id, []).append((float(first["onset"]), gap))
            rows = []
            for event_id in expected_ids:
                row = by_id[event_id]
                onset = float(row["onset_start"])
                duration = float(row["duration"])
                pairs = fallback_pairs.get(event_id, [])
                if not row.get("all_pairs_count"):
                    if not pairs or abs(pairs[-1][0] - onset) > 1e-5 or abs(pairs[-1][1] - duration) > 1e-5:
                        raise ValueError(f"Selected interval does not match the marker stream: {trigger_path}, {event_id}")
                pair_count = int(row["all_pairs_count"]) if row.get("all_pairs_count") else len(pairs)
                is_rest = event_id == 255
                rows.append(
                    {
                        "onset": f"{onset:.6f}",
                        "duration": f"{duration:.6f}",
                        "sample": int(round(onset * 1000)),
                        "trial_type": "rest" if is_rest else "speech",
                        "story_id": "n/a" if is_rest else event_id,
                        "stim_file": "n/a" if is_rest else f"audio/story-{event_id:02d}.mp3",
                        "candidate_pair_count": pair_count,
                        "selection_rule": "last_adjacent_pair_with_duration_ge_150_s",
                    }
                )
            rows.sort(key=lambda row: float(row["onset"]))
            session = f"ses-{day}"
            event_path = (
                release_root
                / subject
                / session
                / "eeg"
                / f"{subject}_{session}_task-audio_events.tsv"
            )
            atomic_write_tsv(event_path, fields, rows)

    atomic_write_json(
        release_root / "task-audio_events.json",
        {
            "onset": {"Description": "Selected interval onset relative to recording start.", "Units": "s"},
            "duration": {"Description": "Selected rest or story interval duration.", "Units": "s"},
            "sample": {"Description": "Onset sample in the 1000-Hz raw recording."},
            "trial_type": {"Description": "Interval type.", "Levels": {"rest": "Eyes-open rest", "speech": "Story speech perception"}},
            "story_id": {"Description": "Story identifier from 1 to 50; n/a for rest."},
            "stim_file": {"Description": "Path relative to the BIDS stimuli directory."},
            "candidate_pair_count": {"Description": "Number of qualifying adjacent marker pairs in the original marker stream."},
            "selection_rule": {"Description": "Deterministic rule used to resolve duplicate or accidental hardware markers."},
        },
    )


def load_manifest(path: Path | None) -> dict[int, dict]:
    if path is None:
        return {}
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return {int(row["story_id"]): row for row in csv.DictReader(handle, delimiter="\t")}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(8 * 1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def copy_stimuli(source_root: Path, release_root: Path, manifest_path: Path | None) -> None:
    manifest = load_manifest(manifest_path)
    audio_dir = release_root / "stimuli" / "audio"
    transcript_dir = release_root / "stimuli" / "transcripts"
    fields = [
        "story_id",
        "session_id",
        "order_within_session",
        "audio_file",
        "transcript_file",
        "tts_platform",
        "voice_name",
        "voice_gender",
        "duration_sec_from_analysis_envelope",
        "chinese_character_count",
        "audio_sha256",
        "transcript_sha256",
    ]
    rows = []
    for story_id in range(1, 51):
        source_audio = source_root / "bids" / "stimuli" / "audio_v7" / f"{story_id}.mp3"
        source_text = source_root / "bids" / "stimuli" / "txt" / f"{story_id}.txt"
        target_audio = audio_dir / f"story-{story_id:02d}.mp3"
        target_text = transcript_dir / f"story-{story_id:02d}.txt"
        hardlink_or_copy(source_audio, target_audio)
        hardlink_or_copy(source_text, target_text)
        text = source_text.read_text(encoding="utf-8-sig")
        cjk_count = sum("\u4e00" <= char <= "\u9fff" for char in text)
        if story_id <= 17:
            session, order = "ses-day1", story_id
        elif story_id <= 33:
            session, order = "ses-day2", story_id - 17
        else:
            session, order = "ses-day3", story_id - 33
        voice_name = "Dapiaoliang" if story_id <= 25 else "Huazai"
        voice_gender = "female" if story_id <= 25 else "male"
        duration = manifest.get(story_id, {}).get("duration_sec_from_100hz_envelope", "n/a")
        rows.append(
            {
                "story_id": story_id,
                "session_id": session,
                "order_within_session": order,
                "audio_file": f"audio/{target_audio.name}",
                "transcript_file": f"transcripts/{target_text.name}",
                "tts_platform": "Guagua Audiobook",
                "voice_name": voice_name,
                "voice_gender": voice_gender,
                "duration_sec_from_analysis_envelope": duration,
                "chinese_character_count": cjk_count,
                "audio_sha256": sha256(source_audio),
                "transcript_sha256": sha256(source_text),
            }
        )
    atomic_write_tsv(release_root / "stimuli" / "stimuli.tsv", fields, rows)
    atomic_write_json(
        release_root / "stimuli" / "stimuli.json",
        {
            "story_id": {"Description": "Story identifier."},
            "session_id": {"Description": "Session in which the story was presented."},
            "order_within_session": {"Description": "Story order within the session."},
            "audio_file": {"Description": "Audio path relative to the stimuli directory."},
            "transcript_file": {"Description": "Transcript path relative to the stimuli directory."},
            "tts_platform": {"Description": "Text-to-speech production platform."},
            "voice_name": {"Description": "Platform voice preset name."},
            "voice_gender": {"Description": "Voice condition used in the experiment."},
            "duration_sec_from_analysis_envelope": {"Description": "Duration of the corresponding 100-Hz analysis envelope.", "Units": "s"},
            "chinese_character_count": {"Description": "Count of U+4E00-U+9FFF characters in the released transcript."},
            "audio_sha256": {"Description": "SHA-256 checksum of the released audio file."},
            "transcript_sha256": {"Description": "SHA-256 checksum of the released transcript."},
        },
    )
    atomic_write_text(
        release_root / "stimuli" / "README",
        """PKUEEG stimulus materials

The directory contains 50 Mandarin story audio files and matching transcripts.
Stories 1-25 use the Guagua Audiobook platform voice preset Dapiaoliang (female),
and stories 26-50 use the platform voice preset Huazai (male). Huazai is recorded
as the platform preset name and does not identify a natural-person narrator.

The narrative text is based on Lin Handa's Three Kingdoms Stories, 1979 edition.
The dataset owner confirmed that the released source text and generated audio may
be distributed with the dataset. Attribution and text integrity are retained.
""",
    )


def copy_behavior(xlsx_path: Path, release_root: Path) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(xlsx_path, data_only=True, read_only=True)
    sheet = workbook["Behavior Data"]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=6, max_row=6))]
    rows = []
    for cells in sheet.iter_rows(min_row=7, max_row=sheet.max_row):
        source = dict(zip(headers, [cell.value for cell in cells]))
        if not source.get("subject"):
            continue
        rows.append(
            {
                "participant_id": source["subject"],
                "session_id": f"ses-{source['day']}",
                "comprehension_accuracy": f"{float(source['correct']):.6f}",
            }
        )
    if len(rows) != 75 or len({(row["participant_id"], row["session_id"]) for row in rows}) != 75:
        raise ValueError("Expected one behavioral record for each of 25 participants and 3 sessions.")
    atomic_write_tsv(
        release_root / "phenotype" / "behavior.tsv",
        ["participant_id", "session_id", "comprehension_accuracy"],
        rows,
    )
    atomic_write_json(
        release_root / "phenotype" / "behavior.json",
        {
            "participant_id": {"Description": "Anonymous participant identifier."},
            "session_id": {"Description": "BIDS session identifier."},
            "comprehension_accuracy": {
                "Description": "Proportion of comprehension questions answered correctly in the session.",
                "Units": "proportion",
                "Minimum": 0,
                "Maximum": 1,
            },
        },
    )


def correct_40hz_npz(source: Path, target: Path, divide_by_million: bool) -> None:
    target.parent.mkdir(parents=True, exist_ok=True)
    if not divide_by_million:
        hardlink_or_copy(source, target)
        return
    with np.load(source, allow_pickle=False) as archive:
        data = archive["eeg_data"].astype(np.float32, copy=True)
        channels = archive["ch_names"]
    data /= np.float32(1_000_000.0)
    tmp = target.with_name(f".{target.name}.tmp-{os.getpid()}")
    with tmp.open("wb") as handle:
        np.savez(handle, eeg_data=data, ch_names=channels)
    os.replace(tmp, target)


def derive_8hz_file(source_40hz: Path, target_8hz: Path) -> tuple[str, tuple[int, int]]:
    with np.load(source_40hz, allow_pickle=False) as archive:
        data = archive["eeg_data"].astype(np.float64, copy=False)
        channels = archive["ch_names"]
    sos = butter(8, 8.0, btype="lowpass", fs=250.0, output="sos")
    filtered = sosfiltfilt(sos, data, axis=1)
    downsampled = resample_poly(filtered, 64, 125, axis=1, window=("kaiser", 5.0), padtype="line")
    downsampled = downsampled.astype(np.float32)
    target_8hz.parent.mkdir(parents=True, exist_ok=True)
    tmp = target_8hz.with_name(f".{target_8hz.name}.tmp-{os.getpid()}")
    with tmp.open("wb") as handle:
        np.savez(handle, eeg_data=downsampled, ch_names=channels)
    os.replace(tmp, target_8hz)
    return str(target_8hz), downsampled.shape


def build_derivatives(source_root: Path, release_root: Path, workers: int) -> None:
    source_40 = source_root / "bids" / "derivatives" / "derivatives_npz" / "formal_filter_40hz"
    target_40 = release_root / "derivatives" / "preprocessed_40hz"
    target_8 = release_root / "derivatives" / "preprocessed_8hz"
    expected_names = {f"story_{index}.npz" for index in range(1, 51)} | {
        "resting-state_1.npz",
        "resting-state_2.npz",
        "resting-state_3.npz",
    }
    for subject in SUBJECTS:
        subject_source = source_40 / subject
        actual_names = {path.name for path in subject_source.glob("*.npz")}
        if actual_names != expected_names:
            raise ValueError(f"Incomplete 40-Hz source for {subject}")
        for source in sorted(subject_source.glob("*.npz")):
            day3 = source.name == "resting-state_3.npz" or (
                source.name.startswith("story_") and int(source.stem.split("_")[1]) >= 34
            )
            correct_40hz_npz(source, target_40 / subject / source.name, day3)

    jobs = []
    with ProcessPoolExecutor(max_workers=max(1, workers)) as executor:
        for subject in SUBJECTS:
            for source in sorted((target_40 / subject).glob("*.npz")):
                jobs.append(executor.submit(derive_8hz_file, source, target_8 / subject / source.name))
        for index, future in enumerate(as_completed(jobs), start=1):
            path, shape = future.result()
            if index % 100 == 0 or index == len(jobs):
                print(f"8-Hz derivative: {index}/{len(jobs)} files; latest={path}; shape={shape}", flush=True)

    shared_description = {
        "BIDSVersion": "1.7.0",
        "DatasetType": "derivative",
        "SourceDatasets": [{"URL": "../../"}],
    }
    description_40 = dict(shared_description)
    description_40.update(
        {
            "Name": "PKUEEG preprocessed EEG 1-40 Hz",
            "GeneratedBy": [
                {
                    "Name": "PKUEEG preprocessing pipeline",
                    "Version": "1.0.0",
                    "Description": "Bad-channel interpolation, 50/100/150-Hz notch filtering, 250-Hz resampling, 1-40-Hz filtering, average reference, and ICA ocular-component removal.",
                }
            ],
        }
    )
    atomic_write_json(target_40 / "dataset_description.json", description_40)
    atomic_write_text(
        target_40 / "README",
        """PKUEEG 1-40 Hz derivative

Each participant directory contains 50 story segments and three eyes-open rest
segments in NPZ format. The eeg_data array is channels by samples, stored as
float32 microvolts. ch_names stores channel labels. The sampling rate is 250 Hz.

The existing preprocessing comprised bad-channel interpolation, 50/100/150-Hz
notch filtering, resampling to 250 Hz, 1-40-Hz filtering, common-average
rereferencing, and ICA component rejection based on correlation with EOG
channels. Day-3 values were divided by 1,000,000 in this release to correct a
documented source-unit interpretation error. No source file was overwritten.
""",
    )
    description_8 = dict(shared_description)
    description_8.update(
        {
            "Name": "PKUEEG preprocessed EEG 1-8 Hz",
            "GeneratedBy": [
                {
                    "Name": "PKUEEG 1-8-Hz release derivation",
                    "Version": "1.0.0",
                    "Description": "Zero-phase eighth-order Butterworth low-pass filtering at 8 Hz followed by polyphase resampling from the corrected 250-Hz derivative to 128 Hz.",
                }
            ],
        }
    )
    atomic_write_json(target_8 / "dataset_description.json", description_8)
    atomic_write_text(
        target_8 / "README",
        """PKUEEG 1-8 Hz derivative

Each participant directory contains 50 story segments and three eyes-open rest
segments in NPZ format. The eeg_data array is channels by samples, stored as
float32 microvolts. ch_names stores channel labels. The sampling rate is 128 Hz.

This derivative was generated uniformly for all 25 participants from the
corrected 1-40-Hz/250-Hz derivative. A zero-phase eighth-order Butterworth
low-pass filter at 8 Hz was applied before polyphase resampling by 64/125.
""",
    )
    atomic_write_json(
        release_root / "derivatives" / "npz_data_dictionary.json",
        {
            "eeg_data": {"Description": "Preprocessed signal array with shape channels by samples.", "Units": "uV", "DataType": "float32"},
            "ch_names": {"Description": "Channel labels in the same order as the first array dimension."},
        },
    )


def write_root_files(root: Path) -> None:
    atomic_write_json(
        root / "dataset_description.json",
        {
            "Name": "PKUEEG",
            "BIDSVersion": "1.7.0",
            "DatasetType": "raw",
            "License": "CC0",
            "Authors": AUTHORS,
            "EthicsApprovals": [
                "Institutional Review Board of Peking University, IRB00001052-25045"
            ],
            "ReferencesAndLinks": ["https://www.gstudios.com.cn/"],
        },
    )
    atomic_write_text(
        root / "README",
        """PKUEEG

PKUEEG contains anonymized EEG from 25 healthy right-handed native Mandarin
speakers during continuous speech perception and eyes-open rest. All participants
completed three recording sessions. Day 1 contains stories 1-17, Day 2 stories
18-33, and Day 3 stories 34-50. Each session includes one approximately 300-s
eyes-open resting interval.

Ethics
------
The protocol was approved by the Institutional Review Board of Peking University
(IRB00001052-25045). All participants provided written informed consent before
data collection. The EEG data were anonymized before release.

Acquisition
-----------
Days 1 and 2 were recorded at 1000 Hz with a NeuroScan SynAmps2 system using 62
scalp EEG channels and two EOG channels. Day 3 was recorded at 1000 Hz with a
NeuSen Wireless EEG/ERP system manufactured by Neuracle, with 59 scalp EEG
channels and five auxiliary channels. Recording system, session order, and story
set are confounded and should not be interpreted as independent experimental
factors. Standard montage labels are provided, but no participant-specific
electrode coordinates were collected.

Events
------
Each events.tsv contains one selected interval per story and one rest interval.
The original hardware marker stream remains in the BrainVision marker file. When
duplicate or accidental markers were present, the released interval used the last
adjacent marker pair whose separation was at least 150 s, matching the segmenting
rule used for the preprocessing derivatives.

Stimuli and behavior
--------------------
Fifty Mandarin story audio files and matching transcripts are included. The
audio was synthesized with the Guagua Audiobook platform. Stories 1-25 use the
Dapiaoliang female preset; stories 26-50 use the Huazai male preset. Session-level
comprehension accuracy is supplied in phenotype/behavior.tsv. Story-level
stimulus derivatives include a 100-Hz speech envelope; uniform 50-Hz Mel, BERT
and Wav2Vec 2.0 representations; 12-layer word-level BERT vectors; and MFA
TextGrid, Pinyin-Hanzi and corrected word timing annotations. Detailed axes,
dimensions and provenance are recorded under stimuli/features and
stimuli/annotations.

Derivatives
-----------
The release includes NPZ derivatives at 1-40 Hz/250 Hz and 1-8 Hz/128 Hz. Signal
values are stored in microvolts. The Day-3 raw BrainVision data and 40-Hz NPZ
segments were corrected by a factor of 1e-6 because the source BDF physical unit
field was empty even though its physical range was expressed in microvolts. The
1-8-Hz derivative was regenerated uniformly from the corrected 40-Hz derivative.

Code
----
The code directory contains a path-independent public preprocessing entry point,
the anonymous manual bad-channel configuration, the release builder, and the
release validator. The `code/technical_validation/` subdirectory contains the
completeness, signal-quality, ISC, TRF, held-out-story prediction and observed
behavior-neural analyses. `code/stimulus_features/` contains the portable
stimulus-derivative packaging and validation workflow. Default inputs and
outputs are release-relative.
""",
    )
    atomic_write_text(
        root / "CHANGES",
        """1.0.0 2026-09-09

- Created the PKUEEG release candidate without modifying source data.
- Added complete ethics, device, session, event, stimulus, and behavior metadata.
- Removed CapTrak-labelled coordinates because individual coordinates were not collected.
- Removed acquisition dates from scans tables.
- Corrected Day-3 raw and preprocessed amplitude scaling by 1e-6.
- Regenerated a complete 1-8-Hz/128-Hz derivative for all participants.
- Included only the 1-8-Hz and 1-40-Hz preprocessing derivatives.
- Added public path-independent preprocessing and release-validation code.
- Added organized technical-validation code and its frozen 100-Hz envelope inputs.
- Added 50-story Mel, time-aligned BERT, Wav2Vec 2.0 layer-9 and word-level BERT features.
- Added MFA TextGrid, Pinyin-Hanzi and corrected word timing annotations.
- Normalized legacy 100-Hz Mel/BERT files for stories 34-43 to the release-wide 50-Hz rate and recorded the operation in the feature manifest.
""",
    )
    atomic_write_text(
        root / "LICENSE",
        """PKUEEG is released under the Creative Commons CC0 1.0 Universal dedication.
See https://creativecommons.org/publicdomain/zero/1.0/ for the legal code.

Attribution to the dataset authors and to Lin Handa for the source narrative text
is requested as a scholarly norm and to preserve authorship information.
""",
    )
    atomic_write_text(
        root / ".bidsignore",
        """code/**
derivatives/**
phenotype/**
stimuli/stimuli.tsv
stimuli/stimuli.json
stimuli/README
stimuli/transcripts/**
stimuli/features/**
stimuli/annotations/**
validation/**
""",
    )


def verify_resumable_raw_scaling(source_raw_root: Path, output_root: Path) -> None:
    """Refuse resume unless every Day-3 binary is an already corrected copy."""
    for subject in SUBJECTS:
        stem = f"{subject}_ses-day3_task-audio_eeg.eeg"
        source_path = source_raw_root / subject / "ses-day3" / "eeg" / stem
        target_path = output_root / subject / "ses-day3" / "eeg" / stem
        if not target_path.exists() or source_path.stat().st_size != target_path.stat().st_size:
            raise ValueError(f"Cannot safely resume: missing or mismatched Day-3 file for {subject}")
        source = np.memmap(source_path, dtype="<f4", mode="r")
        target = np.memmap(target_path, dtype="<f4", mode="r")
        block = min(65_536, source.size)
        starts = sorted({0, max(0, source.size // 2 - block // 2), max(0, source.size - block)})
        source_sample = np.concatenate(
            [np.asarray(source[start : start + block], dtype=np.float64) for start in starts]
        )
        target_sample = np.concatenate(
            [np.asarray(target[start : start + block], dtype=np.float64) for start in starts]
        )
        usable = np.isfinite(source_sample) & np.isfinite(target_sample) & (np.abs(source_sample) > 1e-12)
        if not np.any(usable):
            raise ValueError(f"Cannot safely resume: no usable samples for {subject}")
        ratio = float(np.median(np.abs(target_sample[usable] / source_sample[usable])))
        if not 0.9e-6 <= ratio <= 1.1e-6:
            raise ValueError(f"Cannot safely resume: {subject} Day-3 scaling ratio is {ratio:.6g}")
        del source, target


def main() -> None:
    args = parse_args()
    source_root = args.source_root.resolve()
    output_root = args.output_root.resolve()
    raw_source = source_root / "bids" / "rawdata_bids"
    if args.resume:
        if not output_root.exists():
            raise FileNotFoundError(f"Resume target does not exist: {output_root}")
        verify_resumable_raw_scaling(raw_source, output_root)
        print("Resume safety check passed: all 25 Day-3 raw files already have 1e-6 scaling", flush=True)
    else:
        if output_root.exists():
            raise FileExistsError(f"Output root already exists: {output_root}")
        print(f"Cloning raw BIDS tree into {output_root}", flush=True)
        clone_raw_bids(raw_source, output_root)

        removed = remove_unmeasured_coordinates(output_root)
        print(f"Removed {removed} unmeasured coordinate files", flush=True)

        day3_eeg_files = sorted(output_root.glob("sub-*/ses-day3/eeg/*_eeg.eeg"))
        if len(day3_eeg_files) != 25:
            raise ValueError(f"Expected 25 Day-3 EEG binary files, found {len(day3_eeg_files)}")
        for index, path in enumerate(day3_eeg_files, start=1):
            print(f"Scaling Day-3 raw file {index}/25: {path}", flush=True)
            scale_float32_file_atomic(path, 1e-6)

    update_raw_metadata(output_root)
    write_participants(output_root)
    write_sessions(output_root)
    replace_events_with_intervals(source_root, output_root)
    copy_stimuli(source_root, output_root, args.stimulus_manifest)
    copy_behavior(args.behavior_xlsx, output_root)
    build_derivatives(source_root, output_root, args.workers)
    write_root_files(output_root)

    code_dir = output_root / "code"
    code_dir.mkdir(parents=True, exist_ok=True)
    shutil.copy2(Path(__file__).resolve(), code_dir / Path(__file__).name)
    print(f"Release candidate built at {output_root}", flush=True)


if __name__ == "__main__":
    main()
